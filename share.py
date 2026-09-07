"""
Sharing functionality-check results with a team: a custom message into Slack
with real @-mentions, and the selected tickets as a spreadsheet — an .xlsx
attached in the thread, or a Google Sheet link.

The shape of the feature follows three rules its owner set:

    WYSIWYG     what is shared is exactly the rows the sender had filtered on
                the Functionality Check page — the client sends the ticket
                numbers it is showing, and the sheet is built from those.
    real pings  group mentions go out as <!subteam^ID> and people as <@ID>;
                plain "@handle" text pings nobody, so the compose panel picks
                from the live workspace directory rather than trusting typing.
    no surprises  capability is probed up front (Slack scopes, Drive access)
                and shown in Admin and in the panel, instead of failing on
                Send. Local copies are blocked by the same outward guard as
                every other Slack write.
"""

import asyncio
import io
import logging
from datetime import datetime, timezone

import auth
import funcheck
import gcp
import slack
import vault

logger = logging.getLogger(__name__)

REQUIRED_SCOPES = ("usergroups:read", "files:write")

XLSX_COLUMNS = [
    ("Ticket", "number"),
    ("Title", "title"),
    ("Assignee", "assignee_name"),
    ("Status", "state"),
    ("Date", "fetch_date"),
    ("Functionality (tagged)", "tagged_functionality"),
    ("Functionality OK", "func_ok"),
    ("Suggested functionality", "suggested_functionality"),
    ("New option?", "func_suggestion_new"),
    ("Category (tagged)", "tagged_category"),
    ("Category OK", "cat_ok"),
    ("Suggested category", "suggested_category"),
    ("Finding", "note"),
    ("Pylon link", "link"),
]


def rows_for(month: str, numbers: list[int]) -> list[dict]:
    """The share rows: the month's funcheck view, restricted to `numbers`.

    Numbers not in the month are dropped silently — they cannot be shared
    because there is nothing to share; the caller reports the real count.
    """
    wanted = set()
    for n in numbers or []:
        try:
            wanted.add(int(n))
        except (TypeError, ValueError):
            continue
    data = funcheck.results(month)
    return [t for t in data["tickets"] if t["number"] in wanted]


def build_xlsx(rows: list[dict], month: str) -> bytes:
    """The rows as a real spreadsheet, with a header row and sane widths."""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = f"Functionality check {month}"

    def cell_value(t, key):
        v = t.get(key)
        if key in ("func_ok", "cat_ok"):
            return "" if not t.get("checked") else ("yes" if v else "NO")
        if key in ("func_suggestion_new", "cat_suggestion_new"):
            return "yes" if v else ""
        return v if v is not None else ""

    ws.append([label for label, _ in XLSX_COLUMNS])
    for c in ws[1]:
        c.font = Font(bold=True)
    for t in rows:
        ws.append([cell_value(t, key) for _, key in XLSX_COLUMNS])

    widths = {1: 10, 2: 46, 3: 20, 4: 18, 5: 12, 6: 34, 7: 14, 8: 34,
              9: 12, 10: 34, 11: 12, 12: 34, 13: 60, 14: 40}
    for i, w in widths.items():
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def default_message(month: str, issues: int) -> str:
    """The pre-filled starter the sender edits."""
    pretty = datetime.strptime(month, "%Y-%m").strftime("%B")
    plural = "" if issues == 1 else "s"
    return (f"Functionality tagging review for {pretty} — {issues} "
            f"ticket{plural} need retagging, sheet attached.")


def mention_line(mentions: list[dict]) -> str:
    """Mention tokens from the picker's selections. Unknown types dropped."""
    parts = []
    for m in mentions or []:
        mid = str((m or {}).get("id") or "").strip()
        kind = (m or {}).get("type")
        if not mid:
            continue
        if kind == "group":
            parts.append(f"<!subteam^{mid}>")
        elif kind == "user":
            parts.append(f"<@{mid}>")
    return " ".join(parts)


async def meta() -> dict:
    """Everything the compose panel needs, probed honestly."""
    out: dict = {
        "channel_default": (vault.get_setting("slack_channel") or "").strip(),
        "groups": [],
        "slack_ok": False,
        "slack_message": "",
        # drive_status is sync httpx with real timeouts — off the event loop,
        # or every other request stalls while Google answers.
        "sheet": await asyncio.to_thread(
            gcp.drive_status, vault.get_setting("share_drive_folder_id")),
        "sheet_visibility": vault.get_setting("share_sheet_visibility") or "domain",
    }
    try:
        scopes = await slack.bot_scopes()
        missing = [s for s in REQUIRED_SCOPES if s not in scopes]
        if missing:
            out["slack_message"] = (
                "The Slack app is missing the scope(s) "
                f"{', '.join(missing)} — reinstall it with them to share.")
        else:
            out["slack_ok"] = True
            out["groups"] = await slack.list_usergroups()
    except slack.SlackNotConfigured as e:
        out["slack_message"] = str(e)
    except Exception as e:
        out["slack_message"] = f"Slack probe failed: {str(e)[:200]}"
    return out


async def send(month: str, numbers: list[int], message: str,
               mentions: list[dict], channel: str | None,
               fmt: str, triggered_by: str) -> dict:
    """Build the sheet and deliver message + data to Slack."""
    rows = rows_for(month, numbers)
    if not rows:
        raise ValueError("Nothing to share — the selected tickets were not "
                         "found in this month")
    # The whole action is outward-facing; refuse up front rather than creating
    # an orphan Sheet and then failing on the message.
    if not slack.may_post():
        raise slack.NotTheDeployment(
            "Refusing to share from a non-deployed copy. Set "
            "allow_local_side_effects in Admin if you mean it.")

    channel = (channel or "").strip() \
        or (vault.get_setting("slack_channel") or "").strip()
    if not channel:
        raise ValueError("No Slack channel — configure one in Admin or use "
                         "the override field")

    issues = sum(1 for t in rows
                 if t.get("checked") and not (t["func_ok"] and t["cat_ok"]))
    text = (message or "").strip() or default_message(month, issues)
    tags = mention_line(mentions)
    if tags:
        text = f"{text}\n{tags}"
    stamp = datetime.now(timezone.utc).strftime("%Y%m%d-%H%M")
    title = f"Functionality check {month} ({len(rows)} tickets)"

    result: dict = {"rows": len(rows), "issues": issues,
                    "channel": channel, "format": fmt}
    xlsx = build_xlsx(rows, month)
    if fmt == "sheet":
        # Sync httpx upload+convert, up to two minutes — never on the loop.
        sheet = await asyncio.to_thread(
            gcp.create_sheet_from_xlsx,
            f"Functionality check {month} · {stamp}", xlsx,
            vault.get_setting("share_drive_folder_id"),
            vault.get_setting("share_sheet_visibility") or "domain",
            auth.ALLOWED_DOMAIN)
        posted = await slack._post("chat.postMessage", {
            "channel": channel,
            "text": f"{text}\n<{sheet['link']}|Open the Google Sheet>",
        })
        result.update({"sheet_link": sheet["link"], "ts": posted.get("ts")})
    else:
        posted = await slack.post_with_file(
            channel, text, f"functionality-check-{month}-{stamp}.xlsx",
            xlsx, title)
        result.update({"ts": posted.get("ts"), "file_id": posted.get("file_id")})
    return result
