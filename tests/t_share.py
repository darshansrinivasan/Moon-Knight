"""Sharing funcheck results: WYSIWYG rows, real mention tokens, honest probes.

Pinned because each failure ships wrong things to a real channel: a mention
that renders as dead text instead of pinging, a sheet containing rows the
sender never selected, a share that half-happens (Sheet created, message
blocked), or a capability failure surfacing on Send instead of up front.
"""
import asyncio
import io
import json

import db
import gcp
import share
import slack
import vault

db.init_db()
fails = []
_num = [9500]
M, D = "2026-09", "2026-09-03"
T0 = "2026-09-03T08:00:00+00:00"


def check(name, got, want):
    ok = got == want
    print(f"  {'OK ' if ok else 'FAIL'} {name}: got {got!r}, want {want!r}")
    if not ok:
        fails.append(name)


def ticket(tid, func_ok=1, cat_ok=1):
    _num[0] += 1
    with db.get_conn() as c:
        c.execute(
            "INSERT OR REPLACE INTO tickets (id,number,fetch_date,title,link,"
            "state,assignee_name,custom_fields,fetched_at) VALUES (?,?,?,?,?,?,?,?,?)",
            (tid, _num[0], D, f"Ticket {tid}",
             f"https://app.usepylon.com/issues/{tid}", "closed", "Ann",
             json.dumps({"functionalities": {"value": "HubSpot"},
                         "request_category": {"value": "general_question"}}), T0))
        c.execute(
            "INSERT OR REPLACE INTO func_checks (ticket_id,fetch_date,"
            "tagged_functionality,tagged_category,func_ok,cat_ok,note,"
            "suggested_functionality,suggested_category,func_suggestion_new,"
            "cat_suggestion_new,fingerprint,checked_at)"
            " VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)",
            (tid, D, "HubSpot", "general_question", func_ok, cat_ok,
             "" if func_ok and cat_ok else "wrong tag",
             "" if func_ok else "Integrations : HubSpot", "", 0, 0, "fp", T0))
    return _num[0]


vault.set_raw_setting("qc_rules_json", '{"excluded_states": []}', "t")
n1 = ticket("s1")
n2 = ticket("s2", func_ok=0)
n3 = ticket("s3")

print("=== the sheet holds exactly the selected rows ===")
rows = share.rows_for(M, [n1, n2, 999999, "x", n2])
check("selection validated against the month",
      sorted(r["number"] for r in rows), [n1, n2])
xlsx = share.build_xlsx(rows, M)
from openpyxl import load_workbook
wb = load_workbook(io.BytesIO(xlsx))
ws = wb.active
check("header row present", ws["A1"].value, "Ticket")
check("one data row per selected ticket", ws.max_row, 3)
# funcheck.results orders issues first, so the failing ticket is row 2.
check("verdicts spelled for a non-analyst reader",
      (ws["G2"].value, ws["H2"].value), ("NO", "Integrations : HubSpot"))
try:
    share.rows_for("2026-13", [n1])
    check("bad month rejected", "returned", "ValueError")
except ValueError:
    check("bad month rejected", "ValueError", "ValueError")

print()
print("=== mention tokens are the real thing, never display text ===")
check("group and user tokens",
      share.mention_line([{"type": "group", "id": "S0GROUP"},
                          {"type": "user", "id": "U0PERSON"},
                          {"type": "junk", "id": "X"}, {}]),
      "<!subteam^S0GROUP> <@U0PERSON>")
check("starter message names month and count",
      share.default_message(M, 2),
      "Functionality tagging review for September — 2 tickets need retagging, "
      "sheet attached.")

print()
print("=== capability is probed up front, with the missing piece named ===")


async def scopes_missing():
    return {"chat:write"}


real_scopes = slack.bot_scopes
slack.bot_scopes = scopes_missing
try:
    m = asyncio.run(share.meta())
finally:
    slack.bot_scopes = real_scopes
check("missing scopes are named",
      ("usergroups:read" in m["slack_message"]
       and "files:write" in m["slack_message"], m["slack_ok"]), (True, False))
check("sheet probe explains itself",
      m["sheet"]["ok"] is False and len(m["sheet"]["message"]) > 0, True)

print()
print("=== a share is all-or-nothing behind the outward guard ===")
sent = {}


async def fake_post_with_file(channel, text, filename, content, title):
    sent.update({"channel": channel, "text": text, "filename": filename,
                 "bytes": len(content)})
    return {"ts": "1.2", "channel": "C123", "file_id": "F1"}


real_upload, real_may = slack.post_with_file, slack.may_post
slack.post_with_file = fake_post_with_file
try:
    slack.may_post = lambda: False
    try:
        asyncio.run(share.send(M, [n1], "msg", [], None, "xlsx", "t@x"))
        check("blocked locally before anything is created", "sent", "refused")
    except slack.NotTheDeployment:
        check("blocked locally before anything is created", "refused", "refused")

    slack.may_post = lambda: True
    vault.set_settings({"slack_channel": "#qc-test"}, "t")
    out = asyncio.run(share.send(
        M, [n1, n2], "  ", [{"type": "group", "id": "S9"}], None, "xlsx", "t@x"))
finally:
    slack.post_with_file = real_upload
    slack.may_post = real_may

check("share result counts", (out["rows"], out["issues"]), (2, 1))
check("blank message falls back to the starter, mentions appended",
      sent["text"],
      "Functionality tagging review for September — 1 ticket need retagging, "
      "sheet attached.\n<!subteam^S9>")
check("default channel used when no override", sent["channel"], "#qc-test")
check("a real xlsx went out", sent["bytes"] > 3000, True)
try:
    asyncio.run(share.send(M, [999999], "m", [], None, "xlsx", "t@x"))
    check("empty selection refused", "sent", "ValueError")
except ValueError:
    check("empty selection refused", "ValueError", "ValueError")

print()
print("=== settings are registered, so Admin can actually save them ===")
refused = vault.set_settings({"share_drive_folder_id": "1abc",
                              "share_sheet_visibility": "link"}, "t")
check("keys accepted by the vault registry", refused, [])
check("values persisted",
      (vault.get_setting("share_drive_folder_id"),
       vault.get_setting("share_sheet_visibility")), ("1abc", "link"))

print()
if fails:
    print(f"FAILURES ({len(fails)}): {fails}")
    raise SystemExit(1)
print("ALL SHARE ASSERTIONS PASSED")
